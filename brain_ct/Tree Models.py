# -*- coding: utf-8 -*-
"""
Tree models for worsening on a second brain CT leading to surgery.

XGBoost, Random Forest, Decision Tree and Gradient Boosting.
SMOTE then ADASYN, cross-validation, ROC curves, feature importance,
SHAP and Excel export.
"""

# %% Import packages

import os
import warnings

import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
import shap

from xgboost import XGBClassifier

from sklearn.base import BaseEstimator
from sklearn.model_selection import train_test_split, StratifiedKFold, RandomizedSearchCV
from sklearn.impute import SimpleImputer
from sklearn.preprocessing import StandardScaler

from sklearn.ensemble import RandomForestClassifier, GradientBoostingClassifier
from sklearn.tree import DecisionTreeClassifier

from sklearn.metrics import (
    confusion_matrix,
    accuracy_score,
    recall_score,
    precision_score,
    f1_score,
    roc_auc_score
)

from imblearn.pipeline import Pipeline as ImbPipeline
from imblearn.over_sampling import SMOTE, ADASYN

from openpyxl.drawing.image import Image as OpenpyxlImage

# %% Paths

project_dir = os.path.dirname(os.path.abspath(__file__))
file_path = os.path.join(project_dir, "data", "data1.csv")
output_dir = os.path.join(project_dir, "outputs", "tree_models")


# %% Custom sampler: SMOTE first, then ADASYN if possible

class SMOTEThenADASYN(BaseEstimator):
    """
    Applies SMOTE first, then ADASYN if ADASYN can generate additional positive cases.

    This is used inside the imblearn pipeline, so it is applied only to training folds
    during cross-validation and never to the held-out test set.
    """

    def __init__(
        self,
        positive_label=1,
        smote_target_ratio=0.75,
        final_target_ratio=1.0,
        k_neighbors=5,
        random_state=42
    ):
        self.positive_label = positive_label
        self.smote_target_ratio = smote_target_ratio
        self.final_target_ratio = final_target_ratio
        self.k_neighbors = k_neighbors
        self.random_state = random_state

    def fit_resample(self, X, y):
        y_series = pd.Series(y)
        counts = y_series.value_counts()

        positive_count = counts.get(self.positive_label, 0)
        negative_label = 0 if self.positive_label == 1 else 1
        negative_count = counts.get(negative_label, 0)

        if positive_count < 2 or negative_count < 1:
            return X, y

        X_resampled = X
        y_resampled = y

        # Step 1: SMOTE toward partial balance
        smote_target_positive = int(np.ceil(negative_count * self.smote_target_ratio))

        if smote_target_positive > positive_count:
            safe_k = min(self.k_neighbors, positive_count - 1)

            if safe_k >= 1:
                try:
                    smote = SMOTE(
                        sampling_strategy={self.positive_label: smote_target_positive},
                        k_neighbors=safe_k,
                        random_state=self.random_state
                    )

                    X_resampled, y_resampled = smote.fit_resample(X_resampled, y_resampled)

                except (ValueError, RuntimeError) as e:
                    warnings.warn(f"SMOTE skipped: {e}", RuntimeWarning)

        # Step 2: ADASYN toward final balance, if possible
        y_after_smote = pd.Series(y_resampled)
        counts_after_smote = y_after_smote.value_counts()

        positive_after_smote = counts_after_smote.get(self.positive_label, 0)
        negative_after_smote = counts_after_smote.get(negative_label, 0)

        final_target_positive = int(np.ceil(negative_after_smote * self.final_target_ratio))

        if final_target_positive > positive_after_smote and positive_after_smote >= 2:
            safe_k = min(self.k_neighbors, positive_after_smote - 1)

            if safe_k >= 1:
                try:
                    adasyn = ADASYN(
                        sampling_strategy={self.positive_label: final_target_positive},
                        n_neighbors=safe_k,
                        random_state=self.random_state
                    )

                    X_resampled, y_resampled = adasyn.fit_resample(X_resampled, y_resampled)

                except (ValueError, RuntimeError) as e:
                    # Keep the SMOTE output if ADASYN cannot add samples.
                    warnings.warn(f"ADASYN skipped; keeping SMOTE output: {e}", RuntimeWarning)

        return X_resampled, y_resampled


# %% Load data

if not os.path.isfile(file_path):
    raise FileNotFoundError(f"Dataset not found: {file_path}. Set file_path in the Paths section.")

data = pd.read_csv(file_path)
data.columns = data.columns.str.strip()


# %% Define variables

outcome = "second_ct_findings"

continuous_vars = [
    "age",
    "first_gcs",
    "time_to_first_CT"
]

categorical_vars = [
    "sex",
    "vehicle_trauma",
    "non_vehicle_trauma",
    "epidural",
    "subdural",
    "contusion",
    "skull_base_fx",
    "IVH",
    "SAH",
    "SGH",
    "linear_skull_fx",
    "skull_fx",
    "wound_findings",
    "drugs",
    "impaired_clothing_time"
]

base_predictors = continuous_vars + categorical_vars


# %% Output paths

os.makedirs(output_dir, exist_ok=True)

excel_output_path = os.path.join(
    output_dir,
    "final_tree_models_smote_then_adasyn_shap.xlsx"
)

combined_roc_output_path = os.path.join(
    output_dir,
    "tree_models_combined_test_roc_dense.png"
)

best_model_shap_plot_path = os.path.join(
    output_dir,
    "best_tree_model_shap_summary.png"
)


# %% Check variables exist

all_vars = [outcome] + base_predictors

missing_vars = [var for var in all_vars if var not in data.columns]

if missing_vars:
    print("\nThese variables are not in your dataset:")
    print(missing_vars)

    print("\nAvailable columns are:")
    print(data.columns.tolist())

    raise ValueError("Some variable names are incorrect.")


# %% Prepare dataframe

df = data[[outcome] + base_predictors].copy()

for var in continuous_vars:
    df[var] = pd.to_numeric(df[var], errors="coerce")

if df[outcome].dtype == "object":
    df[outcome] = df[outcome].astype(str).str.strip().str.lower()

    outcome_map = {
        "no": 0,
        "yes": 1,
        "0": 0,
        "1": 1,
        "false": 0,
        "true": 1,
        "negative": 0,
        "positive": 1,
        "absent": 0,
        "present": 1
    }

    df[outcome] = df[outcome].map(outcome_map)

df[outcome] = pd.to_numeric(df[outcome], errors="coerce")
df = df.dropna(subset=[outcome])
if sorted(df[outcome].unique()) != [0, 1]:
    print("\nOutcome values found:")
    print(sorted(df[outcome].unique()))
    raise ValueError("Outcome must be coded exactly as 0 and 1.")

df[outcome] = df[outcome].astype(int)


# %% Feature engineering

ct_finding_vars = [
    "epidural",
    "subdural",
    "contusion",
    "skull_base_fx",
    "IVH",
    "SAH",
    "SGH",
    "linear_skull_fx",
    "skull_fx"
]

for var in ct_finding_vars:
    df[var] = pd.to_numeric(df[var], errors="coerce")

df["gcs_severe"] = np.where(
    df["first_gcs"].isna(),
    np.nan,
    np.where(df["first_gcs"] <= 8, 1, 0)
)

df["gcs_moderate"] = np.where(
    df["first_gcs"].isna(),
    np.nan,
    np.where((df["first_gcs"] >= 9) & (df["first_gcs"] <= 12), 1, 0)
)

df["gcs_mild"] = np.where(
    df["first_gcs"].isna(),
    np.nan,
    np.where(df["first_gcs"] >= 13, 1, 0)
)

# Original variable name; confirm the time unit before interpreting > 6 as hours.
df["delayed_first_ct_6h"] = np.where(
    df["time_to_first_CT"].isna(),
    np.nan,
    np.where(df["time_to_first_CT"] > 6, 1, 0)
)

df["number_of_initial_findings"] = df[ct_finding_vars].sum(axis=1, skipna=True)

hemorrhage_vars = [
    "epidural",
    "subdural",
    "contusion",
    "IVH",
    "SAH"
]

df["any_hemorrhage"] = np.where(
    df[hemorrhage_vars].sum(axis=1, skipna=True) > 0,
    1,
    0
)

skull_fracture_vars = [
    "skull_base_fx",
    "linear_skull_fx",
    "skull_fx"
]

df["any_skull_fracture"] = np.where(
    df[skull_fracture_vars].sum(axis=1, skipna=True) > 0,
    1,
    0
)

engineered_vars = [
    "gcs_severe",
    "gcs_moderate",
    "gcs_mild",
    "delayed_first_ct_6h",
    "number_of_initial_findings",
    "any_hemorrhage",
    "any_skull_fracture"
]


# %% Final predictors

predictors = base_predictors + engineered_vars

X = df[predictors].copy()
y = df[outcome].copy()

if len(categorical_vars) > 0:
    X = pd.get_dummies(
        X,
        columns=categorical_vars,
        drop_first=True,
        dtype=float
    )

for col in X.columns:
    X[col] = pd.to_numeric(X[col], errors="coerce")

X = X.astype(float)


# %% Outcome prevalence

outcome_distribution = pd.DataFrame({
    "Count": y.value_counts(),
    "Percent": y.value_counts(normalize=True) * 100
})

majority_class_accuracy = y.value_counts(normalize=True).max()

print("\nOutcome distribution:")
print(outcome_distribution.round(2))

print("\nMajority-class baseline accuracy:", round(majority_class_accuracy, 4))


# %% Train/test split

X_train, X_test, y_train, y_test = train_test_split(
    X,
    y,
    test_size=0.30,
    random_state=42,
    stratify=y
)

print("\nTraining sample size:", len(X_train))
print("Testing sample size:", len(X_test))

print("\nTraining outcome distribution:")
print(y_train.value_counts())

print("\nTesting outcome distribution:")
print(y_test.value_counts())


# %% Helper functions

def safe_sheet_name(name, suffix):
    for char in "\\/*?:[]":
        name = name.replace(char, " ")
    # Reserve room for the suffix so ROC and importance sheets stay distinct.
    return f"{name[:30 - len(suffix)].rstrip()} {suffix}"


def classification_metrics(y_true, y_prob, y_pred_class):
    tn, fp, fn, tp = confusion_matrix(
        y_true,
        y_pred_class,
        labels=[0, 1]
    ).ravel()

    metrics = {
        "AUC": roc_auc_score(y_true, y_prob),
        "Accuracy": accuracy_score(y_true, y_pred_class),
        "Recall / Sensitivity": recall_score(y_true, y_pred_class, zero_division=0),
        "Specificity": tn / (tn + fp) if (tn + fp) > 0 else np.nan,
        "Precision / PPV": precision_score(y_true, y_pred_class, zero_division=0),
        "NPV": tn / (tn + fn) if (tn + fn) > 0 else np.nan,
        "F1-score": f1_score(y_true, y_pred_class, zero_division=0)
    }

    confusion = {
        "TN": tn,
        "FP": fp,
        "FN": fn,
        "TP": tp
    }

    return metrics, confusion


def dense_roc_curve(y_true, y_prob, n_thresholds=2001):
    """
    Creates ROC data using many fixed cutoffs from 1.001 to -0.001.
    """

    y_true = np.asarray(y_true)
    y_prob = np.asarray(y_prob)

    thresholds = np.linspace(1.001, -0.001, n_thresholds)

    rows = []

    for threshold in thresholds:
        y_pred = np.where(y_prob >= threshold, 1, 0)

        tn, fp, fn, tp = confusion_matrix(
            y_true,
            y_pred,
            labels=[0, 1]
        ).ravel()

        sensitivity = tp / (tp + fn) if (tp + fn) > 0 else np.nan
        specificity = tn / (tn + fp) if (tn + fp) > 0 else np.nan
        false_positive_rate = 1 - specificity

        rows.append({
            "Threshold": threshold,
            "False Positive Rate": false_positive_rate,
            "True Positive Rate": sensitivity,
            "Sensitivity": sensitivity,
            "Specificity": specificity,
            "TN": tn,
            "FP": fp,
            "FN": fn,
            "TP": tp
        })

    return pd.DataFrame(rows)


def get_youden_cutoff(y_true, y_prob, n_thresholds=2001):
    roc_table = dense_roc_curve(
        y_true,
        y_prob,
        n_thresholds=n_thresholds
    )

    roc_table["Youden J"] = (
        roc_table["Sensitivity"] +
        roc_table["Specificity"] -
        1
    )

    best_row = roc_table.sort_values(
        by="Youden J",
        ascending=False
    ).iloc[0]

    return (
        best_row["Threshold"],
        best_row["Sensitivity"],
        best_row["Specificity"],
        best_row["Youden J"]
    )


def get_positive_class_prob(model, X_data):
    return model.predict_proba(X_data)[:, 1]


def get_class1_shap_values(shap_values):
    if isinstance(shap_values, list):
        if len(shap_values) > 1:
            return shap_values[1]
        else:
            return shap_values[0]

    shap_values = np.array(shap_values)

    if shap_values.ndim == 3:
        return shap_values[:, :, 1]

    return shap_values


# %% CV and sampler

cv = StratifiedKFold(
    n_splits=5,
    shuffle=True,
    random_state=42
)

combined_sampler = SMOTEThenADASYN(
    positive_label=1,
    smote_target_ratio=0.75,
    final_target_ratio=1.0,
    k_neighbors=5,
    random_state=42
)


# %% Model factory functions

def make_pipeline(model):
    steps = [
        ("imputer", SimpleImputer(strategy="median")),
        ("scaler", StandardScaler()),
        ("smote_then_adasyn", combined_sampler),
        ("model", model)
    ]

    return ImbPipeline(steps=steps)


model_specs = {
    "XGBoost": {
        "estimator": make_pipeline(
            XGBClassifier(
                objective="binary:logistic",
                eval_metric="logloss",
                random_state=42,
                n_jobs=-1
            )
        ),
        "params": {
            "model__n_estimators": [100, 200, 300, 500],
            "model__learning_rate": [0.01, 0.03, 0.05, 0.1],
            "model__max_depth": [2, 3, 4, 5],
            "model__min_child_weight": [1, 3, 5, 10],
            "model__subsample": [0.6, 0.8, 1.0],
            "model__colsample_bytree": [0.6, 0.8, 1.0],
            "model__gamma": [0, 0.1, 0.5, 1]
        }
    },

    "Random Forest": {
        "estimator": make_pipeline(
            RandomForestClassifier(
                class_weight=None,
                random_state=42,
                n_jobs=-1
            )
        ),
        "params": {
            "model__n_estimators": [200, 500, 1000],
            "model__max_depth": [2, 3, 4, 5, None],
            "model__min_samples_split": [2, 5, 10, 20],
            "model__min_samples_leaf": [1, 2, 5, 10],
            "model__max_features": ["sqrt", "log2", None]
        }
    },

    "Decision Tree": {
        "estimator": make_pipeline(
            DecisionTreeClassifier(
                class_weight=None,
                random_state=42
            )
        ),
        "params": {
            "model__max_depth": [2, 3, 4, 5, 6, None],
            "model__min_samples_split": [2, 5, 10, 20],
            "model__min_samples_leaf": [1, 2, 5, 10],
            "model__criterion": ["gini", "entropy", "log_loss"]
        }
    },

    "Gradient Boosting": {
        "estimator": make_pipeline(
            GradientBoostingClassifier(
                random_state=42
            )
        ),
        "params": {
            "model__n_estimators": [100, 200, 300, 500],
            "model__learning_rate": [0.01, 0.03, 0.05, 0.1],
            "model__max_depth": [2, 3, 4],
            "model__min_samples_split": [2, 5, 10, 20],
            "model__min_samples_leaf": [1, 2, 5, 10],
            "model__subsample": [0.6, 0.8, 1.0]
        }
    }
}


# %% Fit/tune all tree models

models = {}
cv_results_rows = []

n_iter_search = 30

for base_model_name, spec in model_specs.items():

    print("\n======================================")
    print("Tuning:", base_model_name)
    print("Oversampling: SMOTE then ADASYN")
    print("======================================")

    search = RandomizedSearchCV(
        estimator=spec["estimator"],
        param_distributions=spec["params"],
        n_iter=n_iter_search,
        scoring="roc_auc",
        cv=cv,
        random_state=42,
        n_jobs=-1,
        verbose=1,
        error_score=np.nan
    )

    try:
        search.fit(X_train, y_train)

        model_key = f"{base_model_name} - SMOTE then ADASYN"
        models[model_key] = search.best_estimator_

        cv_results_rows.append({
            "Model": base_model_name,
            "Oversampling": "SMOTE then ADASYN",
            "Best CV AUC": search.best_score_,
            "Best Parameters": str(search.best_params_)
        })

    except Exception as e:
        print("\nModel failed:", base_model_name)
        print("Reason:", str(e))

        cv_results_rows.append({
            "Model": base_model_name,
            "Oversampling": "SMOTE then ADASYN",
            "Best CV AUC": np.nan,
            "Best Parameters": "FAILED: " + str(e)
        })


cv_results_table = pd.DataFrame(cv_results_rows)

if not models:
    raise RuntimeError("All model searches failed. Check the errors printed above.")


# %% Evaluate models

metrics_rows = []
confusion_rows = []
cutoff_rows = []
roc_tables = {}
roc_plot_paths = {}
feature_importance_tables = {}
feature_importance_plot_paths = {}
fitted_models = {}

for model_name, model in models.items():

    print("\n======================================")
    print("Evaluating:", model_name)
    print("======================================")

    fitted_models[model_name] = model

    train_pred_prob = get_positive_class_prob(model, X_train)
    test_pred_prob = get_positive_class_prob(model, X_test)

    # Default cutoff
    default_cutoff = 0.50

    train_pred_default = np.where(train_pred_prob >= default_cutoff, 1, 0)
    test_pred_default = np.where(test_pred_prob >= default_cutoff, 1, 0)

    train_metrics_default, train_confusion_default = classification_metrics(
        y_train,
        train_pred_prob,
        train_pred_default
    )

    test_metrics_default, test_confusion_default = classification_metrics(
        y_test,
        test_pred_prob,
        test_pred_default
    )

    # Dense Youden cutoff from training data
    youden_cutoff, youden_sensitivity, youden_specificity, youden_j = get_youden_cutoff(
        y_train,
        train_pred_prob,
        n_thresholds=2001
    )

    train_pred_youden = np.where(train_pred_prob >= youden_cutoff, 1, 0)
    test_pred_youden = np.where(test_pred_prob >= youden_cutoff, 1, 0)

    train_metrics_youden, train_confusion_youden = classification_metrics(
        y_train,
        train_pred_prob,
        train_pred_youden
    )

    test_metrics_youden, test_confusion_youden = classification_metrics(
        y_test,
        test_pred_prob,
        test_pred_youden
    )

    cutoff_rows.append({
        "Model": model_name,
        "Method": "Default cutoff",
        "Cutoff": default_cutoff,
        "Training Sensitivity": np.nan,
        "Training Specificity": np.nan,
        "Youden J": np.nan
    })

    cutoff_rows.append({
        "Model": model_name,
        "Method": "Youden's J dense cutoff",
        "Cutoff": youden_cutoff,
        "Training Sensitivity": youden_sensitivity,
        "Training Specificity": youden_specificity,
        "Youden J": youden_j
    })

    metric_sets = [
        ("Train", "Default cutoff", default_cutoff, train_metrics_default),
        ("Test", "Default cutoff", default_cutoff, test_metrics_default),
        ("Train", "Youden's J dense cutoff", youden_cutoff, train_metrics_youden),
        ("Test", "Youden's J dense cutoff", youden_cutoff, test_metrics_youden)
    ]

    for dataset_name, cutoff_method, cutoff_value, metrics_dict in metric_sets:
        row = {
            "Model": model_name,
            "Dataset": dataset_name,
            "Cutoff Method": cutoff_method,
            "Cutoff": cutoff_value
        }

        row.update(metrics_dict)
        metrics_rows.append(row)

    confusion_sets = [
        ("Train", "Default cutoff", default_cutoff, train_confusion_default),
        ("Test", "Default cutoff", default_cutoff, test_confusion_default),
        ("Train", "Youden's J dense cutoff", youden_cutoff, train_confusion_youden),
        ("Test", "Youden's J dense cutoff", youden_cutoff, test_confusion_youden)
    ]

    for dataset_name, cutoff_method, cutoff_value, confusion_dict in confusion_sets:
        row = {
            "Model": model_name,
            "Dataset": dataset_name,
            "Cutoff Method": cutoff_method,
            "Cutoff": cutoff_value
        }

        row.update(confusion_dict)
        confusion_rows.append(row)

    # Dense ROC data
    train_roc_table = dense_roc_curve(
        y_train,
        train_pred_prob,
        n_thresholds=2001
    )

    test_roc_table = dense_roc_curve(
        y_test,
        test_pred_prob,
        n_thresholds=2001
    )

    train_auc = roc_auc_score(y_train, train_pred_prob)
    test_auc = roc_auc_score(y_test, test_pred_prob)

    roc_train_table = train_roc_table.copy()
    roc_train_table.insert(0, "Dataset", "Train")

    roc_test_table = test_roc_table.copy()
    roc_test_table.insert(0, "Dataset", "Test")

    roc_tables[model_name] = pd.concat(
        [roc_train_table, roc_test_table],
        ignore_index=True
    )

    # Individual ROC plot
    model_roc_path = os.path.join(
        output_dir,
        f"{model_name.replace(' ', '_').replace('-', '_')}_roc_curve_dense.png"
    )

    plt.figure(figsize=(7, 6))

    plt.plot(
        train_roc_table["False Positive Rate"],
        train_roc_table["True Positive Rate"],
        label=f"Train ROC, AUC = {train_auc:.3f}"
    )

    plt.plot(
        test_roc_table["False Positive Rate"],
        test_roc_table["True Positive Rate"],
        label=f"Test ROC, AUC = {test_auc:.3f}"
    )

    plt.plot(
        [0, 1],
        [0, 1],
        linestyle="--",
        label="No discrimination, AUC = 0.500"
    )

    plt.scatter(
        1 - youden_specificity,
        youden_sensitivity,
        label=f"Training Youden cutoff = {youden_cutoff:.3f}"
    )

    plt.xlabel("1 - Specificity / False Positive Rate")
    plt.ylabel("Sensitivity / True Positive Rate")
    plt.title(f"{model_name} ROC Curve")
    plt.legend(loc="lower right", fontsize=8)
    plt.grid(True)

    plt.savefig(model_roc_path, dpi=300, bbox_inches="tight")
    plt.show()
    plt.close()

    roc_plot_paths[model_name] = model_roc_path

    # Built-in feature importance
    underlying_model = model.named_steps["model"]

    if hasattr(underlying_model, "feature_importances_"):
        feature_importance_table = pd.DataFrame({
            "Feature": model.named_steps["imputer"].get_feature_names_out(X_train.columns),
            "Importance": underlying_model.feature_importances_
        }).sort_values(
            by="Importance",
            ascending=False
        )

        feature_importance_tables[model_name] = feature_importance_table

        importance_plot_path = os.path.join(
            output_dir,
            f"{model_name.replace(' ', '_').replace('-', '_')}_feature_importance.png"
        )

        top_features = feature_importance_table.head(20)

        plt.figure(figsize=(8, 6))

        plt.barh(
            top_features["Feature"][::-1],
            top_features["Importance"][::-1]
        )

        plt.xlabel("Importance")
        plt.ylabel("Feature")
        plt.title(f"{model_name} Feature Importance - Top 20")
        plt.tight_layout()

        plt.savefig(
            importance_plot_path,
            dpi=300,
            bbox_inches="tight"
        )

        plt.show()
        plt.close()

        feature_importance_plot_paths[model_name] = importance_plot_path


metrics_table = pd.DataFrame(metrics_rows)
confusion_table = pd.DataFrame(confusion_rows)
cutoff_table = pd.DataFrame(cutoff_rows)


# %% Select best tree model based on test AUC
# This is a descriptive comparison on this test set, not independent model selection.

test_metrics_youden = metrics_table[
    (metrics_table["Dataset"] == "Test") &
    (metrics_table["Cutoff Method"] == "Youden's J dense cutoff")
].copy()

best_tree_model_name = test_metrics_youden.sort_values(
    by="AUC",
    ascending=False
).iloc[0]["Model"]

best_tree_model = models[best_tree_model_name]

print("\nBest tree model based on test AUC:", best_tree_model_name)


# %% SHAP analysis for best tree model

print("\nRunning SHAP for best tree model:", best_tree_model_name)

steps_dict = best_tree_model.named_steps

best_imputer = steps_dict["imputer"]
best_scaler = steps_dict["scaler"]
best_underlying_model = steps_dict["model"]

X_test_imputed = best_imputer.transform(X_test)
X_test_scaled = best_scaler.transform(X_test_imputed)

X_test_processed_df = pd.DataFrame(
    X_test_scaled,
    columns=best_imputer.get_feature_names_out(X_test.columns),
    index=X_test.index
)

max_shap_samples = 300

if len(X_test_processed_df) > max_shap_samples:
    X_shap = X_test_processed_df.sample(
        n=max_shap_samples,
        random_state=42
    )
else:
    X_shap = X_test_processed_df.copy()

explainer = shap.TreeExplainer(best_underlying_model)

raw_shap_values = explainer.shap_values(X_shap)

shap_values_class1 = get_class1_shap_values(raw_shap_values)

shap_importance_table = pd.DataFrame({
    "Feature": X_shap.columns,
    "Mean Absolute SHAP": np.abs(shap_values_class1).mean(axis=0)
}).sort_values(
    by="Mean Absolute SHAP",
    ascending=False
)

print("\nTop SHAP features:")
print(shap_importance_table.head(20).round(4))

plt.figure()

shap.summary_plot(
    shap_values_class1,
    X_shap,
    show=False,
    max_display=20
)

plt.title(f"SHAP Summary Plot - {best_tree_model_name}")
plt.tight_layout()

plt.savefig(
    best_model_shap_plot_path,
    dpi=300,
    bbox_inches="tight"
)

plt.close()

print("\nSHAP plot saved to:", best_model_shap_plot_path)


# %% Combined test ROC curve

plt.figure(figsize=(9, 8))

for model_name, model in fitted_models.items():

    test_pred_prob = get_positive_class_prob(model, X_test)

    test_roc_table = dense_roc_curve(
        y_test,
        test_pred_prob,
        n_thresholds=2001
    )

    test_auc = roc_auc_score(y_test, test_pred_prob)

    plt.plot(
        test_roc_table["False Positive Rate"],
        test_roc_table["True Positive Rate"],
        label=f"{model_name}, AUC = {test_auc:.3f}"
    )

plt.plot(
    [0, 1],
    [0, 1],
    linestyle="--",
    label="No discrimination, AUC = 0.500"
)

plt.xlabel("1 - Specificity / False Positive Rate")
plt.ylabel("Sensitivity / True Positive Rate")
plt.title("Combined Test ROC Curves - Tree Models")
plt.legend(loc="lower right", fontsize=6)
plt.grid(True)

plt.savefig(
    combined_roc_output_path,
    dpi=300,
    bbox_inches="tight"
)

plt.show()
plt.close()


# %% Export to Excel

with pd.ExcelWriter(excel_output_path, engine="openpyxl") as writer:

    outcome_distribution.round(2).to_excel(
        writer,
        sheet_name="Outcome Distribution"
    )

    pd.DataFrame({
        "Majority Class Baseline Accuracy": [majority_class_accuracy]
    }).round(4).to_excel(
        writer,
        sheet_name="Baseline Accuracy",
        index=False
    )

    cv_results_table.to_excel(
        writer,
        sheet_name="CV Results",
        index=False
    )

    metrics_table.round(4).to_excel(
        writer,
        sheet_name="Performance Metrics",
        index=False
    )

    confusion_table.to_excel(
        writer,
        sheet_name="Confusion Matrices",
        index=False
    )

    cutoff_table.round(4).to_excel(
        writer,
        sheet_name="Cutoffs",
        index=False
    )

    for model_name, roc_table in roc_tables.items():
        sheet_name = safe_sheet_name(model_name, "ROC")
        roc_table.round(4).to_excel(
            writer,
            sheet_name=sheet_name,
            index=False
        )

    for model_name, fi_table in feature_importance_tables.items():
        sheet_name = safe_sheet_name(model_name, "Importance")
        fi_table.round(4).to_excel(
            writer,
            sheet_name=sheet_name,
            index=False
        )

    shap_importance_table.round(4).to_excel(
        writer,
        sheet_name="Best Model SHAP",
        index=False
    )

    workbook = writer.book

    for model_name, image_path in roc_plot_paths.items():
        sheet_name = safe_sheet_name(model_name, "ROC Plot")
        sheet = workbook.create_sheet(sheet_name)

        img = OpenpyxlImage(image_path)
        img.width = 600
        img.height = 450

        sheet.add_image(img, "B2")

    for model_name, image_path in feature_importance_plot_paths.items():
        sheet_name = safe_sheet_name(model_name, "FI Plot")
        sheet = workbook.create_sheet(sheet_name)

        img = OpenpyxlImage(image_path)
        img.width = 700
        img.height = 500

        sheet.add_image(img, "B2")

    combined_sheet = workbook.create_sheet("Combined ROC")

    combined_img = OpenpyxlImage(combined_roc_output_path)
    combined_img.width = 650
    combined_img.height = 500

    combined_sheet.add_image(combined_img, "B2")

    shap_sheet = workbook.create_sheet("Best Model SHAP Plot")

    shap_img = OpenpyxlImage(best_model_shap_plot_path)
    shap_img.width = 700
    shap_img.height = 500

    shap_sheet.add_image(shap_img, "B2")


print("\nExcel file saved to:", excel_output_path)
print("Combined ROC image saved to:", combined_roc_output_path)
print("Best model SHAP plot saved to:", best_model_shap_plot_path)
print("Best tree model:", best_tree_model_name)
